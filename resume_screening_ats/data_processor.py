import os
import re
import pandas as pd
import numpy as np
import pdfplumber
from docx import Document

class DataProcessor:
    def __init__(self, csv_path="/home/malik/Desktop/Resume_Screeing-Capstone project/parsed_resumes.csv"):
        self.csv_path = csv_path
        self.df = None
        self.skill_vocabulary = set()
        self.tool_vocabulary = set()
        self.soft_skill_vocabulary = set()
        self.load_and_clean_data()
        self.build_vocabularies()

    def load_and_clean_data(self):
        """Loads the parsed resumes dataset and performs basic cleaning."""
        if not os.path.exists(self.csv_path):
            raise FileNotFoundError(f"Resume CSV not found at {self.csv_path}")
        
        self.df = pd.read_csv(self.csv_path)
        
        # Clean missing values
        self.df['highest_education'] = self.df['highest_education'].fillna('NA')
        self.df['education_field'] = self.df['education_field'].fillna('NA')
        
        # Text fields cleaning
        text_cols = [
            'technical_skills_raw', 'tools_platforms_raw', 'soft_skills_raw',
            'experience_summary', 'project_summary', 'key_achievements'
        ]
        for col in text_cols:
            self.df[col] = self.df[col].fillna('').astype(str).str.strip()

    def build_vocabularies(self):
        """Builds unique vocabularies for skills, tools, and soft skills from the database."""
        if self.df is None:
            return
        
        # Helper to extract comma-separated terms
        def extract_terms(series):
            terms = set()
            for val in series:
                if val:
                    parts = [p.strip().lower() for p in val.split(',')]
                    terms.update([p for p in parts if p])
            return terms

        self.skill_vocabulary = extract_terms(self.df['technical_skills_raw'])
        self.tool_vocabulary = extract_terms(self.df['tools_platforms_raw'])
        self.soft_skill_vocabulary = extract_terms(self.df['soft_skills_raw'])
        
        # Add common industry terms if not present
        extra_skills = {'python', 'machine learning', 'deep learning', 'nlp', 'pytorch', 'tensorflow', 'java', 'sql', 'c++', 'aws', 'azure', 'gcp'}
        self.skill_vocabulary.update(extra_skills)

    @staticmethod
    def extract_text_from_pdf(pdf_path):
        """Extracts text from a PDF file."""
        text = ""
        try:
            with pdfplumber.open(pdf_path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
        except Exception as e:
            print(f"Error reading PDF {pdf_path}: {e}")
        return text

    @staticmethod
    def extract_text_from_docx(docx_path):
        """Extracts text from a DOCX file."""
        text = ""
        try:
            doc = Document(docx_path)
            for para in doc.paragraphs:
                if para.text:
                    text += para.text + "\n"
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text += cell.text + " "
                    text += "\n"
        except Exception as e:
            print(f"Error reading DOCX {docx_path}: {e}")
        return text

    def parse_resume_text(self, text, file_name="uploaded_resume.pdf"):
        """
        Parses raw resume text to extract structured profile attributes
        similar to the columns in the parsed_resumes.csv dataset.
        """
        clean_text = text.replace('\n', ' ')
        clean_text_lower = clean_text.lower()
        
        # 1. Extract Name (Heuristic: usually first few lines)
        lines = [l.strip() for l in text.split('\n') if l.strip()]
        candidate_name = "Candidate"
        if lines:
            # Avoid picking up headings
            for l in lines[:3]:
                if len(l.split()) >= 2 and not any(k in l.lower() for k in ['resume', 'curriculum', 'cv', 'experience', 'education', 'contact', 'email']):
                    candidate_name = l
                    break

        # 2. Extract Skills using vocabulary matching
        matched_tech_skills = []
        for skill in self.skill_vocabulary:
            # Use word boundaries for skill matching to avoid substring issues (e.g. 'c' in 'cloud')
            # except for skills like 'c++', 'c#', '.net'
            pattern = r'\b' + re.escape(skill) + r'\b'
            if skill in ['c++', 'c#', '.net']:
                pattern = re.escape(skill)
            if re.search(pattern, clean_text_lower):
                matched_tech_skills.append(skill.title())
        
        matched_tools = []
        for tool in self.tool_vocabulary:
            pattern = r'\b' + re.escape(tool) + r'\b'
            if re.search(pattern, clean_text_lower):
                matched_tools.append(tool.title())
                
        matched_soft_skills = []
        for sskill in self.soft_skill_vocabulary:
            pattern = r'\b' + re.escape(sskill) + r'\b'
            if re.search(pattern, clean_text_lower):
                matched_soft_skills.append(sskill.title())

        # 3. Extract years of experience
        years_exp = 0
        exp_patterns = [
            r'(\d+)\+?\s*(?:years?|yrs?)\b\s*(?:of\s*)?(?:experience|work|industry)',
            r'(?:experience|work)\s*(?:of\s*)?(\d+)\+?\s*(?:years?|yrs?)',
            r'\b(\d+)\+?\s*years?\b\s*(?:in\s*)?(?:it|industry|role)'
        ]
        for pattern in exp_patterns:
            matches = re.findall(pattern, clean_text_lower)
            if matches:
                # take the maximum found
                found_years = [int(m) for m in matches if m.isdigit()]
                if found_years:
                    years_exp = max(years_exp, max(found_years))
        
        if years_exp == 0:
            # Heuristic: count work experience date ranges e.g. "2018 - 2022" (4 years)
            date_patterns = [
                r'\b(19\d{2}|20\d{2})\s*[-–—]\s*(19\d{2}|20\d{2}|present|current)\b'
            ]
            total_duration = 0
            for pattern in date_patterns:
                matches = re.findall(pattern, clean_text_lower)
                for start, end in matches:
                    start_year = int(start)
                    end_year = 2026 if end in ['present', 'current'] else int(end)
                    duration = end_year - start_year
                    if 0 < duration <= 15:
                        total_duration += duration
            if total_duration > 0:
                years_exp = min(total_duration, 20)  # cap at 20

        if years_exp == 0:
            years_exp = 2  # default fallback if not found

        # 4. Extract Highest Education
        education_levels = {
            'LLM': ['llm', 'master of laws', 'll.m.'],
            'MBA': ['mba', 'master of business administration'],
            'Masters': ['masters', 'm.tech', 'm.sc', 'ms', 'm.s.', 'master of science', 'master of technology', 'post graduate'],
            'Bachelors': ['bachelors', 'b.tech', 'b.sc', 'be', 'b.e.', 'b.a.', 'bba', 'llb', 'll.b.', 'bachelor of science', 'bachelor of engineering', 'bachelor of technology', 'undergraduate']
        }
        highest_edu = 'Bachelors' # Default fallback
        for level, keywords in education_levels.items():
            if any(kw in clean_text_lower for kw in keywords):
                highest_edu = level
                break

        # 5. Extract Education Field
        education_fields = [
            'computer science', 'information technology', 'mathematics', 'statistics', 
            'data science', 'human resources', 'mechanical engineering', 'law',
            'business administration', 'management', 'psychology'
        ]
        edu_field = 'NA'
        for field in education_fields:
            if field in clean_text_lower:
                edu_field = field.title()
                break

        # 6. Extract experience summary, project summary, key achievements
        # Split text into paragraphs and search for section titles
        paragraphs = [p.strip() for p in text.split('\n') if p.strip()]
        
        experience_summary = ""
        project_summary = ""
        key_achievements = ""

        # Simple semantic extraction based on headers
        current_section = None
        for p in paragraphs:
            p_lower = p.lower()
            if any(h in p_lower for h in ['profile', 'summary', 'about me', 'professional summary']) and len(p.split()) < 4:
                current_section = 'summary'
                continue
            elif any(h in p_lower for h in ['experience', 'work history', 'employment', 'career history']) and len(p.split()) < 4:
                current_section = 'experience'
                continue
            elif any(h in p_lower for h in ['project', 'key projects', 'major projects']) and len(p.split()) < 4:
                current_section = 'project'
                continue
            elif any(h in p_lower for h in ['achievement', 'key achievements', 'accomplishments', 'awards']) and len(p.split()) < 4:
                current_section = 'achievements'
                continue
            elif any(h in p_lower for h in ['skills', 'education', 'certifications', 'languages', 'interest']) and len(p.split()) < 4:
                current_section = None
                continue
            
            if current_section == 'summary' and len(experience_summary) < 500:
                experience_summary += p + " "
            elif current_section == 'experience' and len(experience_summary) < 800:
                experience_summary += p + " "
            elif current_section == 'project' and len(project_summary) < 800:
                project_summary += p + " "
            elif current_section == 'achievements' and len(key_achievements) < 500:
                key_achievements += p + " "

        # Fallbacks if section-based extraction failed to find content
        if not experience_summary.strip():
            # Grab first 3 body paragraphs
            body_paras = [p for p in paragraphs if len(p.split()) > 10][:3]
            experience_summary = " ".join(body_paras)
        if not project_summary.strip():
            # Grab next 3 body paragraphs
            body_paras = [p for p in paragraphs if len(p.split()) > 10][3:6]
            project_summary = " ".join(body_paras)
        if not key_achievements.strip():
            key_achievements = "Successfully delivered multiple deliverables on time, meeting all requirements."

        # Truncate to reasonable lengths
        experience_summary = experience_summary.strip()[:1000]
        project_summary = project_summary.strip()[:1000]
        key_achievements = key_achievements.strip()[:1000]

        # 7. Match heuristic flags
        def check_flag(keywords):
            return 1 if any(kw in clean_text_lower for kw in keywords) else 0

        flags = {
            'management_experience_flag': check_flag(['management', 'manager', 'lead', 'director']),
            'people_management_flag': check_flag(['managed team', 'people management', 'led a team', 'team lead', 'direct reports']),
            'project_management_experience_flag': check_flag(['project management', 'pmp', 'scrum master', 'gantt', 'project manager']),
            'agile_scrum_experience_flag': check_flag(['agile', 'scrum', 'sprint', 'jira', 'kanban']),
            'client_facing_experience_flag': check_flag(['client facing', 'stakeholder', 'consulting', 'customer success', 'client interaction']),
            'delivery_lead_experience_flag': check_flag(['delivery lead', 'delivery manager', 'delivery head', 'end-to-end delivery']),
            'cloud_experience_flag': check_flag(['aws', 'azure', 'gcp', 'cloud', 'kubernetes', 'docker', 'devops']),
            'ml_experience_flag': check_flag(['machine learning', 'deep learning', 'nlp', 'computer vision', 'ai', 'neural network', 'pytorch', 'tensorflow']),
            'compliance_experience_flag': check_flag(['compliance', 'regulatory', 'policy', 'audit', 'legal risk', 'gdpr']),
            'enterprise_systems_experience_flag': check_flag(['sap', 'oracle', 'salesforce', 'crm', 'erp', 'enterprise system']),
            'offshore_onsite_model_experience_flag': check_flag(['offshore', 'onsite', 'global delivery', 'cross-border team']),
            'multi_vendor_coordination_flag': check_flag(['vendor coordination', 'vendor management', 'third party', 'multi-vendor']),
            'process_compliance_experience_flag': check_flag(['process compliance', 'standard operating procedure', 'sop', 'iso', 'itil']),
            'documentation_heavy_role_flag': check_flag(['documentation', 'technical writing', 'specifications', 'report writing', 'records']),
            'mentoring_experience_flag': check_flag(['mentoring', 'mentor', 'coaching', 'trained team', 'guidance']),
            'stakeholder_management_experience_flag': check_flag(['stakeholder management', 'stakeholders', 'c-level', 'executive communication']),
        }

        # Profile completeness and keyword density heuristic estimation
        profile_completeness = (
            (1.0 if len(matched_tech_skills) > 3 else 0.5) * 0.3 +
            (1.0 if years_exp > 0 else 0.0) * 0.2 +
            (1.0 if highest_edu != 'NA' else 0.0) * 0.2 +
            (1.0 if len(experience_summary) > 100 else 0.0) * 0.3
        )
        
        keyword_density = min(1.0, (len(matched_tech_skills) + len(matched_tools)) / 15.0)

        # Primary domain & role estimation
        primary_domain = "IT"
        primary_role = "Software Engineer"
        
        domain_keywords = {
            'Data Science': ['data scientist', 'machine learning', 'deep learning', 'nlp', 'data analyst', 'ml engineer', 'pytorch', 'tensorflow', 'pandas'],
            'HR': ['hr', 'human resources', 'talent acquisition', 'recruiter', 'people partner', 'hr business partner'],
            'Legal': ['legal', 'lawyer', 'attorney', 'counsel', 'litigation', 'compliance officer', 'corporate law'],
            'Engineering': ['mechanical engineer', 'process engineer', 'civil engineer', 'electrical engineer'],
            'Management': ['project manager', 'program manager', 'scrum master', 'product manager', 'consultant']
        }
        
        max_domain_matches = 0
        for domain, keywords in domain_keywords.items():
            matches = sum(1 for kw in keywords if kw in clean_text_lower)
            if matches > max_domain_matches:
                max_domain_matches = matches
                primary_domain = domain

        # Role Heuristics
        if primary_domain == 'Data Science':
            if 'analyst' in clean_text_lower:
                primary_role = 'Data Analyst'
            elif 'ml' in clean_text_lower or 'machine learning' in clean_text_lower:
                primary_role = 'ML Engineer'
            elif years_exp < 3:
                primary_role = 'Junior Data Scientist'
            else:
                primary_role = 'Data Scientist'
        elif primary_domain == 'HR':
            if 'manager' in clean_text_lower:
                primary_role = 'HR Manager'
            elif 'acquisition' in clean_text_lower or 'recruiter' in clean_text_lower:
                primary_role = 'Talent Acquisition Specialist'
            elif 'partner' in clean_text_lower:
                primary_role = 'HR Business Partner'
            else:
                primary_role = 'HR Executive'
        elif primary_domain == 'Legal':
            if 'compliance' in clean_text_lower:
                primary_role = 'Compliance Officer'
            elif 'associate' in clean_text_lower:
                primary_role = 'Legal Associate'
            elif 'litigation' in clean_text_lower:
                primary_role = 'Litigation Lawyer'
            else:
                primary_role = 'Corporate Lawyer'
        elif primary_domain == 'Engineering':
            if 'process' in clean_text_lower:
                primary_role = 'Process Engineer'
            elif years_exp < 3:
                primary_role = 'Junior Mechanical Engineer'
            else:
                primary_role = 'Mechanical Engineer'
        elif primary_domain == 'Management':
            if 'senior' in clean_text_lower or years_exp > 8:
                primary_role = 'Senior Project Manager'
            elif 'consultant' in clean_text_lower:
                primary_role = 'Management Consultant'
            elif 'coordinator' in clean_text_lower:
                primary_role = 'Project Coordinator'
            else:
                primary_role = 'Project Manager'
        else: # IT Domain
            if 'cloud' in clean_text_lower:
                primary_role = 'Cloud Engineer'
            elif 'devops' in clean_text_lower:
                primary_role = 'DevOps Engineer'
            elif 'full stack' in clean_text_lower:
                primary_role = 'Full Stack Developer'
            elif 'senior' in clean_text_lower:
                primary_role = 'Senior Software Engineer'
            elif years_exp < 3:
                primary_role = 'Junior Software Engineer'
            else:
                primary_role = 'Software Engineer'

        profile = {
            'candidate_id': "CND_UPLOADED",
            'resume_file_name': file_name,
            'candidate_name': candidate_name,
            'primary_domain': primary_domain,
            'primary_role': primary_role,
            'years_experience': years_exp,
            'highest_education': highest_edu,
            'education_field': edu_field,
            'institution_tier': 'Tier-2', # default average
            'technical_skills_raw': ", ".join(matched_tech_skills),
            'tools_platforms_raw': ", ".join(matched_tools),
            'soft_skills_raw': ", ".join(matched_soft_skills),
            'experience_summary': experience_summary,
            'project_summary': project_summary,
            'key_achievements': key_achievements,
            'resume_length_words': len(text.split()),
            'keyword_density_score': round(keyword_density, 2),
            'profile_completeness_score': round(profile_completeness, 2),
        }
        
        # Merge flags
        profile.update(flags)
        
        return profile

    def parse_job_description(self, jd_text):
        """Parses a job description to extract target requirements (skills, years exp, domain, etc.)"""
        jd_lower = jd_text.lower()
        
        # 1. Extract required years of experience
        years_req = 0
        exp_patterns = [
            r'(\d+)\+?\s*(?:years?|yrs?)\b\s*(?:of\s*)?(?:experience|work|industry)',
            r'(?:experience|work|requirement)\s*(?:of\s*)?(\d+)\+?\s*(?:years?|yrs?)',
            r'\b(\d+)\+?\s*years?\s*(?:required|preferred|experience)\b'
        ]
        for pattern in exp_patterns:
            matches = re.findall(pattern, jd_lower)
            if matches:
                found_years = [int(m) for m in matches if m.isdigit()]
                if found_years:
                    years_req = max(years_req, max(found_years))
        
        if years_req == 0:
            years_req = 3 # default standard assumption if not specified

        # 2. Extract domain
        jd_domain = "IT"
        domain_keywords = {
            'Data Science': ['data science', 'machine learning', 'deep learning', 'nlp', 'data scientist', 'data analyst', 'statistics', 'ai'],
            'HR': ['hr', 'human resources', 'recruitment', 'talent acquisition', 'people partner'],
            'Legal': ['legal', 'compliance', 'law', 'lawyer', 'counsel', 'litigation'],
            'Engineering': ['mechanical', 'civil', 'electrical', 'process engineer'],
            'Management': ['project management', 'scrum master', 'agile', 'product manager', 'consultant']
        }
        max_matches = 0
        for domain, keywords in domain_keywords.items():
            matches = sum(1 for kw in keywords if kw in jd_lower)
            if matches > max_matches:
                max_matches = matches
                jd_domain = domain

        # 3. Extract required skills/tools/soft skills from JDs using the vocabulary
        req_tech_skills = []
        for skill in self.skill_vocabulary:
            pattern = r'\b' + re.escape(skill) + r'\b'
            if skill in ['c++', 'c#', '.net']:
                pattern = re.escape(skill)
            if re.search(pattern, jd_lower):
                req_tech_skills.append(skill)
        
        req_tools = []
        for tool in self.tool_vocabulary:
            pattern = r'\b' + re.escape(tool) + r'\b'
            if re.search(pattern, jd_lower):
                req_tools.append(tool)

        req_soft_skills = []
        for sskill in self.soft_skill_vocabulary:
            pattern = r'\b' + re.escape(sskill) + r'\b'
            if re.search(pattern, jd_lower):
                req_soft_skills.append(sskill)

        # 4. Extract required education
        highest_edu_req = "Bachelors" # Standard base
        if any(kw in jd_lower for kw in ['masters', 'm.tech', 'm.sc', 'master of', 'postgraduate', 'ms']):
            highest_edu_req = "Masters"
        if any(kw in jd_lower for kw in ['mba', 'master of business']):
            highest_edu_req = "MBA"
        if any(kw in jd_lower for kw in ['llm', 'master of laws']):
            highest_edu_req = "LLM"

        return {
            'domain': jd_domain,
            'years_experience': years_req,
            'technical_skills': req_tech_skills,
            'tools_platforms': req_tools,
            'soft_skills': req_soft_skills,
            'highest_education': highest_edu_req
        }
